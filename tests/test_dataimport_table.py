"""Tests für dataimport/table.py."""

from datetime import datetime

import pytest

from tapesmith.dataimport.table import load_table, read_clipboard_text, read_csv, read_lines, read_xlsx


def test_read_csv_semicolon_cp1252_umlaute(tmp_path):
    text = "Name;Ort\r\nMüller;Köln\r\nSchröder;Bonn\r\n"
    path = tmp_path / "personen.csv"
    path.write_bytes(text.encode("cp1252"))
    table = read_csv(path)
    assert table.headers == ("Name", "Ort")
    assert table.rows == (("Müller", "Köln"), ("Schröder", "Bonn"))
    assert table.source == "personen.csv"


def test_read_csv_comma_utf8_bom(tmp_path):
    text = "a,b\r\n1,2\r\n3,4\r\n"
    path = tmp_path / "daten.csv"
    path.write_bytes(text.encode("utf-8-sig"))
    table = read_csv(path)
    assert table.headers == ("a", "b")
    assert table.rows == (("1", "2"), ("3", "4"))


def test_read_csv_tab_delimiter_from_text():
    text = "a\tb\r\n1\tc\r\n"
    table = read_csv(text, name="tab.txt")
    assert table.headers == ("a", "b")
    assert table.rows == (("1", "c"),)


def test_read_csv_has_header_false_explicit():
    text = "1,2\r\n3,4\r\n"
    table = read_csv(text, has_header=False)
    assert table.headers == ("Spalte 1", "Spalte 2")
    assert table.rows == (("1", "2"), ("3", "4"))


def test_read_csv_numeric_first_row_no_header():
    text = "1,2\r\n3,4\r\n"
    table = read_csv(text)
    assert table.headers == ("Spalte 1", "Spalte 2")
    assert table.rows == (("1", "2"), ("3", "4"))


def test_read_csv_short_row_padded():
    text = "a,b\r\n1\r\n"
    table = read_csv(text, delimiter=",")
    assert table.rows == (("1", ""),)


def test_read_csv_long_row_raises():
    text = "a,b\r\n1,2,3\r\n"
    with pytest.raises(ValueError):
        read_csv(text, delimiter=",")


def test_read_csv_empty_raises():
    with pytest.raises(ValueError, match="Keine Daten"):
        read_csv("")


def test_read_xlsx(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Host", "Slot", "SN", "Datum"])
    ws.append(["pmx10", 3, 274913, datetime(2026, 9, 27)])
    ws.append(["srv2", None, 12.0, None])
    path = tmp_path / "geraete.xlsx"
    wb.save(path)

    table = read_xlsx(path)
    assert table.headers == ("Host", "Slot", "SN", "Datum")
    assert table.rows[0] == ("pmx10", "3", "274913", "27.09.2026")
    assert table.rows[1] == ("srv2", "", "12", "")


def test_read_xlsx_missing_sheet_raises(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["a"])
    path = tmp_path / "x.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="Fehlt"):
        read_xlsx(path, sheet="Fehlt")


def test_read_clipboard_text_with_tabs():
    table = read_clipboard_text("Host\tSN\r\npmx10\t274913\r\n")
    assert table.headers == ("Host", "SN")
    assert table.rows == (("pmx10", "274913"),)
    assert table.source == "Zwischenablage"


def test_read_clipboard_text_without_tabs():
    table = read_clipboard_text("nur ein Wert")
    assert table.headers == ("Spalte 1",)
    assert table.rows == (("nur ein Wert",),)


def test_read_lines():
    table = read_lines("a\n\nb\n")
    assert table.headers == ("Wert",)
    assert table.rows == (("a",), ("b",))
    assert table.source == "Liste"


def test_load_table_by_extension(tmp_path):
    csv_path = tmp_path / "a.csv"
    csv_path.write_text("a,b\r\n1,2\r\n", encoding="utf-8")
    table = load_table(csv_path)
    assert table.headers == ("a", "b")

    tsv_path = tmp_path / "a.tsv"
    tsv_path.write_text("a\tb\r\n1\t2\r\n", encoding="utf-8")
    table_tsv = load_table(tsv_path)
    assert table_tsv.headers == ("a", "b")

    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["a", "b"])
    wb.active.append(["1", "2"])
    xlsx_path = tmp_path / "a.xlsx"
    wb.save(xlsx_path)
    table_xlsx = load_table(xlsx_path)
    assert table_xlsx.headers == ("a", "b")

    pdf_path = tmp_path / "a.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    with pytest.raises(ValueError):
        load_table(pdf_path)
