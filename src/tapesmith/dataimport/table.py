"""Tabellen aus CSV, XLSX, Zwischenablage (Tab-getrennt) und Zeilenlisten lesen."""

import csv
import io
import warnings
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from tapesmith.i18n import _t

_DELIMITER_CANDIDATES = ";,\t|"


@dataclass(frozen=True)
class Table:
    headers: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]
    source: str


def _decode_bytes(data: bytes, encoding: str | None) -> str:
    if encoding is not None:
        return data.decode(encoding)
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252")


def _sniff_delimiter(sample: str) -> str:
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=_DELIMITER_CANDIDATES)
        return dialect.delimiter
    except csv.Error:
        return ";"


def _is_numeric_cell(cell: str) -> bool:
    try:
        float(cell)
        return True
    except ValueError:
        return False


def _looks_like_header(first_row: list[str], has_more_rows: bool) -> bool:
    if not has_more_rows or not first_row:
        return False
    if any(cell == "" for cell in first_row):
        return False
    if all(_is_numeric_cell(cell) for cell in first_row):
        return False
    if len(set(first_row)) != len(first_row):
        return False
    return True


def _build_table(raw_rows: list[list[str]], has_header: bool | None, source: str) -> Table:
    while raw_rows and all(cell == "" for cell in raw_rows[-1]):
        raw_rows.pop()
    if not raw_rows:
        raise ValueError(_t("Keine Daten gefunden"))
    if has_header is None:
        has_header = _looks_like_header(raw_rows[0], len(raw_rows) > 1)
    if has_header:
        headers = tuple(raw_rows[0])
        data_rows = raw_rows[1:]
    else:
        width = max(len(r) for r in raw_rows)
        headers = tuple(_t("Spalte {value}", value=i + 1) for i in range(width))
        data_rows = raw_rows
    width = len(headers)
    rows = []
    for r in data_rows:
        if len(r) < width:
            r = list(r) + [""] * (width - len(r))
        elif len(r) > width:
            raise ValueError(_t("Zeile hat mehr Spalten ({count}) als die Kopfzeile ({width})", count=len(r), width=width))
        rows.append(tuple(r))
    return Table(headers=headers, rows=tuple(rows), source=source)


def read_csv(
    source: Path | str,
    *,
    encoding: str | None = None,
    delimiter: str | None = None,
    has_header: bool | None = None,
    name: str | None = None,
) -> Table:
    if isinstance(source, str):
        text = source
        default_name = name or _t("Text")
    else:
        path = Path(source)
        text = _decode_bytes(path.read_bytes(), encoding)
        default_name = name or path.name
    if delimiter is None:
        delimiter = _sniff_delimiter(text[:4096])
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_rows = [[cell.strip() for cell in row] for row in reader]
    return _build_table(raw_rows, has_header, default_name)


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    return str(value)


def read_xlsx(path: Path, *, sheet: str | None = None, has_header: bool | None = None) -> Table:
    import openpyxl

    path = Path(path)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is None:
            ws = wb.worksheets[0]
        else:
            if sheet not in wb.sheetnames:
                raise ValueError(_t("Blatt '{sheet}' gibt es nicht (vorhanden: {items})", sheet=sheet, items=', '.join(wb.sheetnames)))
            ws = wb[sheet]
        raw_rows = [[_cell_to_str(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return _build_table(raw_rows, has_header, path.name)


def read_clipboard_text(text: str, *, has_header: bool | None = None) -> Table:
    raw_rows = [[cell.strip() for cell in line.split("\t")] for line in text.splitlines()]
    return _build_table(raw_rows, has_header, _t("Zwischenablage"))


def read_lines(text: str) -> Table:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return Table(headers=(_t("Wert"),), rows=tuple((line,) for line in lines), source="Liste")


def load_table(path: Path, **kw) -> Table:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".tsv":
        kw.setdefault("delimiter", "\t")
        return read_csv(path, **kw)
    if suffix in (".csv", ".txt"):
        return read_csv(path, **kw)
    if suffix in (".xlsx", ".xlsm"):
        return read_xlsx(path, **kw)
    raise ValueError(_t("Dateityp '{suffix}' wird nicht unterstützt", suffix=suffix))
